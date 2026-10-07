"""Ordinary seven-slide PPTX for native relationship host checks.

build FILE: create an untouched baseline; use a copy in PowerPoint.
verify FILE [bound|detached|final]: inspect saved OOXML independently of VBA.
build-batch FILE / verify-batch FILE: one parent and a separate child group.
No test macros or source-only diagnostic presentations are created.
"""
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import xml.etree.ElementTree as ET

spec = importlib.util.spec_from_file_location('native_host_fixture', Path(__file__).with_name('native-host-fixture.py'))
host_fixture = importlib.util.module_from_spec(spec)
spec.loader.exec_module(host_fixture)


def build(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Cm, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Cm(30), Cm(18)
    tag_number = 0

    def tags(slide, shape, values):
        nonlocal tag_number
        tag_number += 1
        root = ET.Element('{%s}tagLst' % host_fixture.NS['p'])
        for key, value in values.items():
            ET.SubElement(root, '{%s}tag' % host_fixture.NS['p'], name=key.upper(), val=str(value))
        part = Part(PackURI('/ppt/tags/tag%d.xml' % tag_number),
                    'application/vnd.openxmlformats-officedocument.presentationml.tags+xml',
                    slide.part.package, ET.tostring(root, encoding='utf-8'))
        relationship = slide.part.relate_to(part, host_fixture.NS['r'] + '/tags')
        nv = shape.element.xpath('./p:nvSpPr/p:nvPr | ./p:nvGrpSpPr/p:nvPr')[0]
        data = OxmlElement('p:custDataLst')
        tag = OxmlElement('p:tags')
        tag.set('{%s}id' % host_fixture.NS['r'], relationship)
        data.append(tag)
        nv.append(data)

    def rounded(container, name, x, y, w, h, frame=False):
        shape = container.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
        shape.name = name
        shape.text = name
        shape.adjustments[0] = 0.15
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(229, 236, 247) if frame else RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(80, 100, 125)
        shape.text_frame.paragraphs[0].font.size = Pt(14)
        shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(25, 40, 60)
        return shape

    def slide():
        return presentation.slides.add_slide(presentation.slide_layouts[6])

    s = slide()
    p = rounded(s.shapes, 'ParentBox', 1, 2, 13, 9, True)
    tags(s, p, {'custom': 'KeepParentCase', 'radiusLock_v1': '1.35'})
    for name, x in [('ChildA', 2), ('ChildB', 6), ('ChildC', 10)]:
        child = rounded(s.shapes, name, x, 4, 3, 4)
        values = {'custom': 'Keep' + name + 'Case'}
        if name == 'ChildB':
            values.update({'radiusLockStrict_v1': '1', 'radiusLock_v1': '0.45'})
        tags(s, child, values)
    rounded(s.shapes, 'OtherParent', 17, 2, 11, 9, True)
    rounded(s.shapes, 'OtherChild', 19, 4, 6, 4)
    s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Cm(2), Cm(13), Cm(5), Cm(2)).name = 'UntouchedArrow'

    s = slide()
    outer = s.shapes.add_group_shape()
    outer.name = 'OuterRelationGroup'
    p = rounded(outer.shapes, 'GroupParent', 1, 2, 14, 9, True)
    inner = outer.shapes.add_group_shape()
    inner.name = 'InnerRelationGroup'
    a = rounded(inner.shapes, 'NestedA', 3, 4, 4, 4)
    b = rounded(inner.shapes, 'NestedB', 9, 4, 4, 4)
    arrow = inner.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Cm(6), Cm(10), Cm(4), Cm(1))
    arrow.name = 'NestedArrow'
    outer.element.recalculate_extents()
    tags(s, p, {'radiusRelation_v1': 'G07', 'radiusRelationRole_v1': 'P', 'custom': 'KeepParentCase'})
    tags(s, a, {'radiusRelation_v1': 'G07', 'radiusRelationRole_v1': 'C1', 'custom': 'KeepNestedACase'})
    tags(s, b, {'radiusRelation_v1': 'G07', 'radiusRelationRole_v1': 'C2', 'radiusLockStrict_v1': '1', 'custom': 'KeepNestedBCase'})
    tags(s, inner, {'customInner': 'KeepInnerCase'})
    tags(s, outer, {'customOuter': 'KeepOuterCase'})
    outer.width = int(outer.width * 1.15)
    outer.height = int(outer.height * 0.9)

    s = slide()
    p = rounded(s.shapes, 'LegacyParent', 2, 2, 9, 6, True)
    c = rounded(s.shapes, 'LegacyChild', 4, 4, 5, 3)
    tags(s, p, {'layoutParent_v1': '  ' + json.dumps({'rows': 1, 'cols': 1, 'padding': 0.35, 'gutter': 0.2, 'linkRMode': 'same', 'extra': 'KeepCase', 'childIds': [str(c.shape_id)]}) + '  ', 'custom': 'KeepLegacyParentCase'})
    tags(s, c, {'layoutChild_v1': str(p.shape_id), 'custom': 'KeepLegacyChildCase'})
    orphan = rounded(s.shapes, 'OrphanChild', 18, 4, 5, 3)
    tags(s, orphan, {'layoutChild_v1': '42', 'custom': 'KeepOrphanCase', 'radiusLockStrict_v1': '1'})

    s = slide()
    for name, x in [('DuplicateParentA', 2), ('DuplicateParentB', 12)]:
        shape = rounded(s.shapes, name, x, 2, 7, 6, True)
        tags(s, shape, {'radiusRelation_v1': 'G09', 'radiusRelationRole_v1': 'P'})

    s = slide()
    rounded(s.shapes, 'SlideSwitchTarget', 2, 2, 7, 6)
    s = slide()
    s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Cm(2), Cm(2), Cm(7), Cm(3)).name = 'OnlyArrow'
    s = slide()
    rounded(s.shapes, 'RadiusSmoke', 2, 2, 6, 4)
    presentation.save(path)


def build_batch(path):
    from pptx import Presentation
    with tempfile.TemporaryDirectory() as directory:
        baseline = Path(directory) / 'baseline.pptx'
        build(baseline)
        presentation = Presentation(baseline)
    slide = presentation.slides[0]
    members = [shape for shape in slide.shapes
               if shape.name in ('ChildA', 'ChildB', 'ChildC', 'UntouchedArrow')]
    slide.shapes.add_group_shape(members).name = 'BatchChildren'
    for shape in list(slide.shapes):
        if shape.name in ('OtherParent', 'OtherChild'):
            shape.element.getparent().remove(shape.element)
    for slide_id in list(presentation.slides._sldIdLst)[1:]:
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    presentation.save(path)


def verify_batch(path):
    with tempfile.TemporaryDirectory() as directory:
        baseline = Path(directory) / 'baseline.pptx'
        build_batch(baseline)
        before = host_fixture.inspect(baseline, 1)['1']
    after = host_fixture.inspect(path, 1)['1']
    assert before.keys() == after.keys(), 'Batch object names/count changed'
    roles = {'ParentBox': 'P', 'ChildA': 'C1', 'ChildB': 'C2', 'ChildC': 'C3'}
    for name, original in before.items():
        actual = after[name]
        expected = dict(original['tags'])
        if name in roles:
            expected.update({'RADIUSRELATION_V1': 'G01', 'RADIUSRELATIONROLE_V1': roles[name]})
        assert actual['tags'] == expected, f'Unexpected batch tags: {name}'
        assert actual['group'] == original['group'] and actual['parent'] == original['parent'], name
        assert all(abs(x-y) < 0.0001 for x, y in zip(actual['boxCm'], original['boxCm'])), name
        if not original['group']:
            assert actual['id'] == original['id'], name
        if 'radiusCm' in original:
            assert abs(actual['radiusCm'] - original['radiusCm']) < 0.0001, name
    return {'stage': 'batch', 'file': str(path), 'checks': [
        'One Ribbon bind creates a parent and three ordered children',
        'Strict protection, fixed radius, custom tags and radii preserved',
        'Child group hierarchy, leaf IDs and geometry preserved; arrow untagged',
    ], 'savedState': {'1': after}}


def verify(path, stage='final'):
    with tempfile.TemporaryDirectory() as directory:
        baseline = Path(directory) / 'baseline.pptx'
        build(baseline)
        before = host_fixture.inspect(baseline)
    after = host_fixture.inspect(path)
    expected = {slide: {name: dict(shape['tags']) for name, shape in shapes.items()}
                for slide, shapes in before.items()}

    def native(slide, name, key, role):
        expected[slide][name].update({'RADIUSRELATION_V1': key, 'RADIUSRELATIONROLE_V1': role})

    native('1', 'ParentBox', 'G01', 'P')
    for name, role in [('ChildA', 'C1'), ('ChildB', 'C2')]:
        native('1', name, 'G01', role)
    if stage == 'bound':
        native('1', 'ChildC', 'G01', 'C3')
    if stage in ('bound', 'detached'):
        native('1', 'OtherParent', 'G02', 'P')
        native('1', 'OtherChild', 'G02', 'C1')
    if stage == 'final':
        for name, role in [('GroupParent', 'P'), ('NestedA', 'C1'), ('NestedB', 'C2')]:
            native('2', name, 'G01', role)
        expected['3']['LegacyParent'].pop('LAYOUTPARENT_V1')
        expected['3']['LegacyChild'].pop('LAYOUTCHILD_V1')
        expected['3']['OrphanChild'].pop('LAYOUTCHILD_V1')
    for slide, shapes in before.items():
        assert shapes.keys() == after[slide].keys(), f'Object names/count changed on slide {slide}'
        for name, original in shapes.items():
            actual = after[slide][name]
            context = f'{slide}/{name}'
            assert actual['group'] == original['group'], f'Group type changed: {context}'
            assert actual['parent'] == original['parent'], f'Group hierarchy changed: {context}'
            assert all(abs(x-y) < 0.0001 for x,y in zip(actual['boxCm'], original['boxCm'])), f'Geometry changed: {context}'
            if not original['group']:
                assert actual['id'] == original['id'], f'Leaf identity changed: {context}'
            assert actual['tags'] == expected[slide][name], f'Unexpected tags: {context}: {actual["tags"]}'
            if 'radiusCm' in original:
                target = 0.3 if stage == 'final' and slide == '7' else original['radiusCm']
                assert abs(actual['radiusCm'] - target) < 0.0001, f'Radius changed: {context}'
    return {'stage': stage, 'file': str(path), 'checks': [
        'Native relationship tags match the explicit parent and ordered children',
        'Existing custom, fixed-radius and strict-protection tags preserved',
        'Geometry, names, leaf IDs and nested group hierarchy preserved',
        'No preview badge shapes in the saved source presentation',
        'Duplicate metadata slide and non-rounded shape unchanged',
    ], 'savedState': after}


if __name__ == '__main__':
    command, file = sys.argv[1:3]
    if command == 'build':
        build(Path(file))
        print(file)
    elif command == 'build-batch':
        build_batch(Path(file))
        print(file)
    elif command == 'verify-batch':
        print(json.dumps(verify_batch(file), indent=2, ensure_ascii=False))
    elif command == 'inspect':
        print(json.dumps(host_fixture.inspect(file), indent=2, ensure_ascii=False))
    elif command == 'verify':
        print(json.dumps(verify(file, sys.argv[3] if len(sys.argv) > 3 else 'final'), indent=2, ensure_ascii=False))
    else:
        raise SystemExit('Use build, build-batch, inspect, verify or verify-batch')
