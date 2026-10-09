"""Ordinary PPTX for radius edits beside an unrelated broken relationship.

Build a disposable presentation, apply the 0.30cm preset only to
IndependentRadius on slide 4, then save and verify. No test macros are used.
The duplicate parents deliberately remain broken and must not be repaired by
the radius command. Other slides and all unrelated objects remain unchanged.
"""
import argparse
import importlib.util
import json
from pathlib import Path
import tempfile
import xml.etree.ElementTree as ET
import zipfile


spec = importlib.util.spec_from_file_location(
    'relations_fixture', Path(__file__).with_name('native-relations-fixture.py'))
relations = importlib.util.module_from_spec(spec)
spec.loader.exec_module(relations)
host = relations.host_fixture
TARGET = 'IndependentRadius'


def build(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Cm

    relations.build(path)
    presentation = Presentation(path)
    shape = presentation.slides[3].shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Cm(2), Cm(11), Cm(5), Cm(3))
    shape.name = shape.text = TARGET
    shape.adjustments[0] = 0.2
    presentation.save(path)

    # A configured duplicate parent triggers the old full-slide dependency.
    # Edit the existing tag part rather than introducing a second tags record.
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info)) for info in archive.infolist()]
    modified = 0
    with zipfile.ZipFile(path, 'w') as archive:
        for info, data in entries:
            if info.filename.startswith('ppt/tags/') and info.filename.endswith('.xml'):
                root = ET.fromstring(data)
                tags = {tag.get('name').upper(): tag.get('val') for tag in root}
                if tags.get('RADIUSRELATION_V1') == 'G09' and tags.get('RADIUSRELATIONROLE_V1') == 'P':
                    ET.SubElement(root, '{%s}tag' % host.NS['p'],
                                  name='RADIUSRELATIONLAYOUT_V1',
                                  val='1|1|1|0.300000|0.200000|same|1')
                    data = ET.tostring(root, encoding='utf-8')
                    modified += 1
            archive.writestr(info, data)
    if modified != 2:
        raise ValueError('Expected two deliberately duplicated parents.')


def verify(path):
    with tempfile.TemporaryDirectory(prefix='radius-independent-baseline-') as directory:
        baseline = Path(directory) / 'baseline.pptx'
        build(baseline)
        before = host.inspect(baseline)
    with zipfile.ZipFile(path) as archive:
        presentation = ET.fromstring(archive.read('ppt/presentation.xml'))
        slides = presentation.find('p:sldIdLst', host.NS)
        if slides is None or len(slides) != len(before):
            raise AssertionError('Slide count changed.')
    after = host.inspect(path)
    for slide, originals in before.items():
        saved = after[slide]
        if saved.keys() != originals.keys():
            raise AssertionError('Shape names/count changed on slide ' + slide)
        for name, original in originals.items():
            actual = saved[name]
            context = slide + '/' + name
            for key in ('id', 'group', 'parent', 'tags'):
                if actual[key] != original[key]:
                    raise AssertionError(context + ': changed ' + key)
            if any(abs(a - b) > 0.0001 for a, b in zip(actual['boxCm'], original['boxCm'])):
                raise AssertionError(context + ': changed geometry')
            if 'radiusCm' in original:
                expected = 0.3 if slide == '4' and name == TARGET else original['radiusCm']
                if abs(actual['radiusCm'] - expected) > 0.0001:
                    raise AssertionError(context + ': unexpected radius')
    return {'file': str(path), 'inspectedSlides': len(after), 'checks': [
        'IndependentRadius is 0.30cm despite unrelated duplicate configured parents',
        'Broken relationship tags remain intact; no silent repair',
        'Other radii, geometry, hierarchy, IDs and tags remain unchanged',
    ]}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=('build', 'inspect', 'verify'))
    parser.add_argument('file', type=Path)
    args = parser.parse_args()
    if args.command == 'build':
        build(args.file)
    else:
        result = host.inspect(args.file) if args.command == 'inspect' else verify(args.file)
        print(json.dumps(result, ensure_ascii=False, indent=2))
