"""Ordinary nine-slide PPTX; validate saved native layout/link outcomes.

build FILE creates a baseline. verify FILE checks the final host protocol.
inspect FILE prints saved OOXML for intermediate snapshots.
"""
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import xml.etree.ElementTree as ET

spec = importlib.util.spec_from_file_location('host_fixture', Path(__file__).with_name('native-host-fixture.py'))
host = importlib.util.module_from_spec(spec)
spec.loader.exec_module(host)
PT = 360000 / 12700
TOL = 0.0002
SETTINGS = 'RADIUSRELATIONLAYOUT_V1'
BASELINE = 'RADIUSRELATIONBASELINE_V1'


def build(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Cm, Pt
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Cm(30), Cm(18)
    tag_index = 0

    def tags(slide, shape, values):
        nonlocal tag_index
        tag_index += 1
        root = ET.Element('{%s}tagLst' % host.NS['p'])
        for key, value in values.items():
            ET.SubElement(root, '{%s}tag' % host.NS['p'], name=key.upper(), val=str(value))
        part = Part(PackURI('/ppt/tags/tag%d.xml' % tag_index),
                    'application/vnd.openxmlformats-officedocument.presentationml.tags+xml',
                    slide.part.package, ET.tostring(root, encoding='utf-8'))
        rel = slide.part.relate_to(part, host.NS['r'] + '/tags')
        nv = shape.element.xpath('./p:nvSpPr/p:nvPr | ./p:nvGrpSpPr/p:nvPr')[0]
        data = OxmlElement('p:custDataLst')
        tag = OxmlElement('p:tags')
        tag.set('{%s}id' % host.NS['r'], rel)
        data.append(tag)
        nv.append(data)

    def rounded(container, name, x, y, w, h, radius, parent=False):
        shape = container.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
        shape.name, shape.text = name, name
        shape.adjustments[0] = radius / min(w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(205, 220, 245) if parent else RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(45, 75, 115)
        shape.text_frame.paragraphs[0].font.size = Pt(13)
        return shape

    specs = [(12, 8, 1.2, 4), (10, 8, 1, 1), (8, 6, 0.9, 2),
             (9, 6, 0.9, 2), (4, 3, 0.45, 2), (14, 9, 1.35, 4),
             (9, 6, 0.9, 1), (12, 7, 1, 3), (10, 6, 0.9, 1)]
    for index, (width, height, radius, count) in enumerate(specs, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        container = slide.shapes
        if index == 6:
            outer = container.add_group_shape()
            outer.name = 'OuterLayoutGroup'
            container = outer.shapes
        parent = rounded(container, 'ParentBox', 2, 2, width, height, radius, True)
        if index == 6:
            inner = container.add_group_shape()
            inner.name = 'InnerLayoutGroup'
            container = inner.shapes
        children = []
        for i in range(count):
            child = rounded(container, 'Child' + chr(65+i), 3 + (i % 2) * 5,
                            4 + (i // 2) * 3, 3, 2, 0.3 if i % 2 == 0 else 0.4)
            children.append(child)
            values = {'radiusRelation_v1': 'G01', 'radiusRelationRole_v1': 'C%d' % (i+1),
                      'custom': 'KeepChild%dCase' % (i+1), 'radiusLock_v1': '0.3' if i % 2 == 0 else '0.4'}
            if index == 4 and i == count-1:
                values['radiusLockStrict_v1'] = '1'
            tags(slide, child, values)
        values = {'radiusRelation_v1': 'G01', 'radiusRelationRole_v1': 'P', 'custom': 'KeepParentCase'}
        if index in (1, 6, 8, 9):
            values['radiusLock_v1'] = str(radius)
        if index == 9:
            values['radiusLockStrict_v1'] = '1'
        if index in (4, 7):
            columns = count
            values['radiusRelationLayout_v1'] = '1|1|%d|0.300000|0.200000|same|%d' % (columns, index == 4)
            box = [2*PT, 2*PT, width*PT, height*PT]
            members = ''.join('%d:%d,' % (child.shape_id, i+1) for i, child in enumerate(children))
            values['radiusRelationBaseline_v1'] = '|'.join('%.6f' % v for v in box+[radius]) + '|' + members
        tags(slide, parent, values)
        if index == 6:
            arrow = container.add_shape(MSO_SHAPE.RIGHT_ARROW, Cm(6), Cm(10), Cm(4), Cm(1))
            arrow.name = 'UntouchedArrow'
            outer.element.recalculate_extents()
            tags(slide, inner, {'customInner': 'KeepInnerCase'})
            tags(slide, outer, {'customOuter': 'KeepOuterCase'})
            outer.width = int(outer.width * 1.15)
            outer.height = int(outer.height * 0.9)
    presentation.save(path)


def inspect(path):
    return host.inspect(path, 9)


def close(a, b, context):
    assert abs(a-b) < TOL, f'{context}: {a} != {b}'


def verify(path):
    with tempfile.TemporaryDirectory() as directory:
        baseline_path = Path(directory)/'baseline.pptx'
        build(baseline_path)
        before = inspect(baseline_path)
    after = inspect(path)
    configs = {'1': (2, 2, .5, .3, 'subtract', True), '2': (1, 1, .5, .2, 'subtract', True),
               '3': (1, 2, .3, .2, 'off', True), '6': (2, 2, .4, .2, 'subtract', True),
               '8': (1, 3, .3, .2, 'same', True), '9': (1, 1, .3, .2, 'same', True)}
    for slide, shapes in before.items():
        actual = after[slide]
        assert shapes.keys() == actual.keys(), f'Shape names/count changed: slide {slide}'
        for name, original in shapes.items():
            saved = actual[name]
            context = slide + '/' + name
            assert saved['group'] == original['group'] and saved['parent'] == original['parent'], context
            if not original['group']:
                assert saved['id'] == original['id'], context
            expected_tags = dict(original['tags'])
            if slide in configs and name == 'ParentBox':
                fields = saved['tags'][SETTINGS].split('|')
                rows, cols, pad, gap, mode, automatic = configs[slide]
                assert fields[0] == '1' and tuple(map(int, fields[1:3])) == (rows, cols), context
                close(float(fields[3]), pad, context)
                close(float(fields[4]), gap, context)
                assert fields[5:] == [mode, str(int(automatic))], context
                assert len(saved['tags'][BASELINE].split('|')) == 6, context
                expected_tags[SETTINGS] = saved['tags'][SETTINGS]
                expected_tags[BASELINE] = saved['tags'][BASELINE]
            if 'RADIUSLOCK_V1' in expected_tags and ((slide in ('1', '6', '8') and name == 'ParentBox') or (slide in ('1', '2', '6', '8', '9') and name.startswith('Child'))):
                expected_tags['RADIUSLOCK_V1'] = saved['tags']['RADIUSLOCK_V1']
                close(float(saved['tags']['RADIUSLOCK_V1']), saved['radiusCm'], context+' fixed')
            assert saved['tags'] == expected_tags, f'Unexpected tags: {context}: {saved["tags"]}'
            if slide not in configs or name == 'UntouchedArrow':
                if name != 'ParentBox' or slide not in ('4', '7'):
                    for x, y in zip(saved['boxCm'], original['boxCm']): close(x, y, context+' untouched geometry')
                if 'radiusCm' in original and name != 'ParentBox': close(saved['radiusCm'], original['radiusCm'], context+' untouched radius')
        parent = actual['ParentBox']
        if slide in ('2', '3', '4', '7'):
            assert parent['boxCm'][2] > shapes['ParentBox']['boxCm'][2]+.5, f'No host resize: {slide}'
            if slide == '2':
                for x,y in zip(parent['boxCm'], [3,3,12,10]): close(x,y, '2 parent moved/resized')
        else:
            for x,y in zip(parent['boxCm'], shapes['ParentBox']['boxCm']): close(x,y, slide+' parent box')
        if slide in ('1', '3', '6', '7', '8'):
            close(parent['radiusCm'], {'1': .8, '3': .3, '6': 1, '7': .5, '8': .5}[slide], slide+' parent radius')
        elif slide in ('2', '4', '5', '9'):
            close(parent['fraction'], shapes['ParentBox']['fraction'], slide+' parent fraction preserved')
        if slide in configs:
            rows, cols, pad, gap, mode, automatic = configs[slide]
            left, top, width, height = parent['boxCm']
            child_width = (width-2*pad-(cols-1)*gap)/cols
            child_height = (height-2*pad-(rows-1)*gap)/rows
            children = sorted(name for name in actual if name.startswith('Child'))
            for i, name in enumerate(children):
                child = actual[name]
                desired = [left+pad+(i%cols)*(child_width+gap), top+pad+(i//cols)*(child_height+gap), child_width, child_height]
                for x,y in zip(child['boxCm'], desired): close(x,y, slide+'/'+name+' grid')
                if mode == 'off':
                    close(child['fraction'], shapes[name]['fraction'], slide+'/'+name+' off fraction')
                else:
                    target = parent['radiusCm'] if mode == 'same' else max(0, parent['radiusCm']-pad)
                    close(child['radiusCm'], min(target, min(child_width,child_height)/2), slide+'/'+name+' linked radius')
    return {'file': str(path), 'passedSlides': 9, 'checks': [
        '2x2 grid and parent radius / inset links; fixed-radius tags updated',
        'Parent resized and moved in PowerPoint; constant cm padding and linked radius follow automatically',
        'R-link off preserves adjustment fractions while geometry follows the parent',
        'Late strict child rejects the entire auto layout and disables parent radius edits',
        'Impossible padding rejects all geometry/radius/settings writes',
        'Scaled nested group retains hierarchy, names, custom tags, leaf IDs and non-rounded arrow',
        'Automatic linking disabled: children remain unchanged after parent resize/radius edits',
        'Saved configuration survives reopening and still links parent/child radius',
        'Protected parent can lay out unprotected children without changing parent radius/protection tags'
    ], 'savedState': after}


def verify_controls(path):
    """Saved-state checks for layout arrows; separate from the nine-page link protocol."""
    with tempfile.TemporaryDirectory() as directory:
        baseline_path = Path(directory) / 'baseline.pptx'
        build(baseline_path)
        before = inspect(baseline_path)
    after = inspect(path)
    configs = {'1': (2, 2, .5, .3), '6': (2, 2, .4, .2)}
    for slide, shapes in before.items():
        actual = after[slide]
        if slide not in configs:
            assert actual == shapes, f'Unexpected saved write on untouched slide {slide}'
            continue
        assert actual.keys() == shapes.keys(), f'Shape names/count changed: slide {slide}'
        rows, cols, padding, gap = configs[slide]
        parent = actual['ParentBox']
        fields = parent['tags'][SETTINGS].split('|')
        assert fields[:3] == ['1', str(rows), str(cols)] and fields[5:] == ['same', '1']
        close(float(fields[3]), padding, slide + ' padding')
        close(float(fields[4]), gap, slide + ' gap')
        assert len(parent['tags'][BASELINE].split('|')) == 6
        for x, y in zip(parent['boxCm'], shapes['ParentBox']['boxCm']):
            close(x, y, slide + ' parent geometry')
        close(parent['fraction'], shapes['ParentBox']['fraction'], slide + ' parent fraction')
        left, top, width, height = parent['boxCm']
        child_width = (width - 2 * padding - (cols - 1) * gap) / cols
        child_height = (height - 2 * padding - (rows - 1) * gap) / rows
        radius = min(parent['radiusCm'], min(child_width, child_height) / 2)
        children = sorted(name for name in actual if name.startswith('Child'))
        for name, original in shapes.items():
            saved = actual[name]
            context = slide + '/' + name
            assert saved['group'] == original['group'] and saved['parent'] == original['parent'], context
            if not original['group']:
                assert saved['id'] == original['id'], context + ' leaf ID'
            expected_tags = dict(original['tags'])
            if name == 'ParentBox':
                expected_tags[SETTINGS] = saved['tags'][SETTINGS]
                expected_tags[BASELINE] = saved['tags'][BASELINE]
            elif name in children:
                i = children.index(name)
                desired = [left + padding + (i % cols) * (child_width + gap),
                           top + padding + (i // cols) * (child_height + gap),
                           child_width, child_height]
                for x, y in zip(saved['boxCm'], desired):
                    close(x, y, context + ' grid')
                close(saved['radiusCm'], radius, context + ' same radius')
                if 'RADIUSLOCK_V1' in expected_tags:
                    expected_tags['RADIUSLOCK_V1'] = saved['tags']['RADIUSLOCK_V1']
                    close(float(saved['tags']['RADIUSLOCK_V1']), radius, context + ' fixed radius')
            elif not original['group']:
                for x, y in zip(saved['boxCm'], original['boxCm']):
                    close(x, y, context + ' untouched geometry')
            assert saved['tags'] == expected_tags, context + ' custom/relation/protection tags'
    return {'file': str(path), 'passedSlides': 9, 'appliedSlides': [1, 6], 'checks': [
        'Arrow-configured 2x2 grids, centimeter spacing, same radius and fixed-radius tags',
        'Scaled nested hierarchy, group names/custom tags, leaf IDs and non-rounded arrow preserved',
        'Parent geometry, adjustment and original tags preserved',
        'Single-child, strict-child and other untouched slides retain all baseline data'
    ], 'savedState': after}


if __name__ == '__main__':
    command, file = sys.argv[1:3]
    if command == 'build':
        build(Path(file))
        print(file)
    elif command == 'inspect':
        print(json.dumps(inspect(file), indent=2, ensure_ascii=False))
    elif command == 'verify':
        print(json.dumps(verify(file), indent=2, ensure_ascii=False))
    elif command == 'verify-controls':
        print(json.dumps(verify_controls(file), indent=2, ensure_ascii=False))
    else:
        raise SystemExit('Use build, inspect, verify or verify-controls')
